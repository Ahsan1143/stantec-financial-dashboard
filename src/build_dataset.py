"""Build the Stantec dashboard dataset.

Inputs
  * Stantec_DCF_Model.xlsx (Muhammad Ahsan's forecast & DCF model): historical
    statements FY2023-2025, forecast FY2026-2030, WACC, DCF, comps, sensitivity.
  * Company-reported figures not in the model (segment net revenue, adjusted
    EBITDA, net debt to adjusted EBITDA, DSO, backlog) - typed below with sources.

Outputs
  data/stantec_dashboard_data.json          -> web dashboard data
  dashboard/index.html                      -> web dashboard (template + data inlined)
  powerbi/Stantec_Dashboard_Data.xlsx       -> Power BI data source (star schema)

Usage:  python src/build_dataset.py path/to/Stantec_DCF_Model.xlsx
"""
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MODEL = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "Stantec_DCF_Model.xlsx"

# ---------------------------------------------------------------- reported figures (C$ millions)
SRC_FY25 = "Stantec Q4/FY2025 results release, Feb 25, 2026"
SRC_FY24P = "Stantec Q4/FY2024 results presentation, Feb 2025"
SRC_MDA23 = "Stantec 2023 MD&A (SEC Form 40-F exhibit 99.2)"
SEGMENTS = [  # year, segment, net revenue, source
    (2023, "Canada", 1246.3, SRC_MDA23), (2023, "United States", 2684.1, SRC_MDA23), (2023, "Global", 1135.8, SRC_MDA23),
    (2024, "Canada", 1427.0, SRC_FY25), (2024, "United States", 3040.7, SRC_FY25), (2024, "Global", 1398.9, SRC_FY25),
    (2025, "Canada", 1546.5, SRC_FY25), (2025, "United States", 3369.4, SRC_FY25), (2025, "Global", 1578.9, SRC_FY25),
]
REPORTED = {  # year: (adjusted EBITDA, net debt / adj. EBITDA, DSO days, backlog)
    2023: (831.0, 1.0, 77, 6305.9),
    2024: (980.3, 1.2, 77, 7823.9),
    2025: (1143.7, 1.3, 69, 8570.2),
}
REPORTED_SRC = {2023: SRC_MDA23, 2024: f"{SRC_FY25} (adj. EBITDA, backlog); {SRC_FY24P} (leverage, DSO)", 2025: SRC_FY25}


def main() -> None:
    wb = load_workbook(MODEL, data_only=True)
    H, F, A, D, C, W = (wb[s] for s in ["Historicals", "Forecast", "Assumptions", "DCF", "Comps", "WACC"])

    # ---------------- actuals FY2023-2025 from the model's Historicals tab
    rows = {r[0].value: [c.value for c in r[1:4]] for r in H.iter_rows(min_row=4, max_row=36) if r[0].value}
    fact = []
    for i, yr in enumerate([2023, 2024, 2025]):
        g = lambda k: rows[k][i]
        adj, lev, dso, backlog = REPORTED[yr]
        fact.append({
            "Year": yr, "Scenario": "Actual",
            "GrossRevenue": g("Gross revenue"), "SubconsultantExpenses": g("Subconsultant and other direct expenses"),
            "NetRevenue": g("Net revenue"), "DirectPayroll": g("Direct payroll costs"), "ProjectMargin": g("Project margin"),
            "AdminMarketing": g("Administrative and marketing"), "DeprPropertyEquipment": g("Depreciation of property and equipment"),
            "DeprLeaseAssets": g("Depreciation of lease assets"), "NetImpairment": g("Net impairment"),
            "AmortIntangibles": g("Amortization of intangible assets"), "EBIT": g("EBIT (operating profit)"),
            "NetInterestExpense": g("Net interest and finance expense"), "OtherIncome": g("Other income"),
            "IncomeTaxes": g("Income taxes"), "NetIncome": g("Net income"),
            "OperatingCashFlow": g("Cash flows from operating activities"), "Capex": g("Capex (property & equipment + intangibles)"),
            "Acquisitions": g("Business acquisitions"), "DividendsPaid": g("Dividends paid"),
            "Cash": g("Cash and cash equivalents"), "Receivables": g("Trade and other receivables"),
            "ContractAssets": g("Contract assets (work done, not yet billed)"), "Payables": g("Trade and other payables"),
            "ContractLiabilities": g("Contract liabilities (billed, not yet done)"),
            "TotalDebt": g("Total debt (excl. lease liabilities)"), "ShareholdersEquity": g("Total shareholders' equity"),
            "AdjEBITDA_Reported": adj, "NetDebtToAdjEBITDA_Reported": lev, "DSO_Reported": dso, "Backlog": backlog,
        })

    # ---------------- forecast FY2026-2030 from the model's Forecast and DCF tabs
    fr = {r[0].value: [c.value for c in r[1:7]] for r in F.iter_rows(min_row=4, max_row=19) if r[0].value}
    dfac = [D.cell(row=14, column=c).value for c in range(3, 8)]
    pv = [D.cell(row=15, column=c).value for c in range(3, 8)]
    for i, yr in enumerate(range(2026, 2031)):
        k = i + 1  # column offset (0 = FY2025A)
        nr, pm = fr["Net revenue"][k], fr["Project margin"][k]
        fact.append({
            "Year": yr, "Scenario": "Forecast", "NetRevenue": nr, "DirectPayroll": nr - pm, "ProjectMargin": pm,
            "AdminMarketing": fr["Administrative and marketing"][k], "DeprPropertyEquipment": fr["Depreciation of property and equipment"][k],
            "DeprLeaseAssets": fr["Depreciation of lease assets"][k], "NetImpairment": 0.0,
            "AmortIntangibles": fr["Amortization of intangible assets"][k], "EBIT": fr["EBIT"][k],
            "TaxOnEBIT": fr["Tax on EBIT"][k], "NOPAT": fr["NOPAT (after-tax operating profit)"][k],
            "Capex": fr["Less: capex"][k], "IncreaseInNWC": fr["Less: increase in working capital"][k],
            "UnleveredFCF": fr["Unlevered free cash flow"][k], "DiscountFactor": dfac[i], "PV_UnleveredFCF": pv[i],
        })
    fact = pd.DataFrame(fact)
    num = fact.select_dtypes("number").columns.drop("Year")
    fact[num] = fact[num].astype(float).round(2)

    seg = pd.DataFrame(SEGMENTS, columns=["Year", "Segment", "NetRevenue", "Source"])

    # ---------------- valuation outputs
    val = [
        ("WACC", W["B18"].value, "%", "WACC tab"), ("Cost of equity", W["B7"].value, "%", "WACC tab"),
        ("Terminal growth", A["B15"].value, "%", "Assumptions tab"),
        ("Sum of PV of free cash flows", D["B18"].value, "C$mm", "DCF tab"),
        ("Terminal value (2030)", D["B19"].value, "C$mm", "DCF tab"), ("PV of terminal value", D["B20"].value, "C$mm", "DCF tab"),
        ("Enterprise value", D["B21"].value, "C$mm", "DCF tab"), ("Debt (excl. leases)", D["B22"].value, "C$mm", "DCF tab"),
        ("Cash", D["B23"].value, "C$mm", "DCF tab"), ("Equity value", D["B24"].value, "C$mm", "DCF tab"),
        ("Diluted shares (millions)", D["B8"].value, "millions", "Assumptions tab"),
        ("DCF value per share", D["B25"].value, "C$", "DCF tab"), ("Share price", D["B26"].value, "C$", "Assumptions tab (October 2026)"),
        ("DCF upside", D["B27"].value, "%", "DCF tab"), ("Terminal value % of EV", D["B28"].value, "%", "DCF tab"),
        ("Peer median EV/EBITDA", C["C8"].value, "x", "Comps tab"), ("Stantec EV/EBITDA", C["C10"].value, "x", "Comps tab"),
        ("Comps value per share", C["C18"].value, "C$", "Comps tab"),
    ]
    valuation = pd.DataFrame(val, columns=["Metric", "Value", "Unit", "Source"])
    sens = []
    gs = [D.cell(row=30, column=c).value for c in range(3, 6)]
    for r in range(31, 37):
        w = D.cell(row=r, column=2).value
        for j, gval in enumerate(gs):
            sens.append({"WACC": w, "TerminalGrowth": gval, "ValuePerShare": round(D.cell(row=r, column=3 + j).value, 2)})
    sens = pd.DataFrame(sens)
    comps = pd.DataFrame([{"Company": C.cell(row=r, column=1).value, "Ticker": C.cell(row=r, column=2).value,
                           "EV_EBITDA": C.cell(row=r, column=3).value} for r in (4, 5, 6, 7, 10)])
    bridge = pd.DataFrame([
        (1, "PV of FCF 2026-30", D["B18"].value), (2, "PV of terminal value", D["B20"].value),
        (3, "Less: debt", -D["B22"].value), (4, "Plus: cash", D["B23"].value),
    ], columns=["StepOrder", "Step", "Amount"])
    years = pd.DataFrame({"Year": range(2023, 2031)})
    years["YearLabel"] = "FY" + years["Year"].astype(str)
    years["Scenario"] = ["Actual"] * 3 + ["Forecast"] * 5
    years["ChartLabel"] = years["YearLabel"] + years["Scenario"].map({"Actual": "A", "Forecast": "E"})

    # ---------------- write JSON for the web dashboard
    (ROOT / "data").mkdir(exist_ok=True)
    out = {
        "fact": json.loads(fact.to_json(orient="records")), "segments": json.loads(seg.drop(columns="Source").to_json(orient="records")),
        "valuation": {m: v for m, v, *_ in val}, "sensitivity": json.loads(sens.to_json(orient="records")),
        "comps": json.loads(comps.to_json(orient="records")),
    }
    (ROOT / "data" / "stantec_dashboard_data.json").write_text(json.dumps(out, indent=1))
    tpl = (ROOT / "dashboard" / "template.html").read_text()
    (ROOT / "dashboard" / "index.html").write_text(tpl.replace("/*__DATA__*/null", json.dumps(out, separators=(",", ":"))))

    # ---------------- write the Power BI data workbook
    sources = pd.DataFrame([
        ("FactFinancials", "FY2023-2025 statements", "Stantec_DCF_Model.xlsx, Historicals tab (from Stantec FY2025 results release and 2024 Annual Report)"),
        ("FactFinancials", "FY2026-2030 forecast, FCF, discount factors", "Stantec_DCF_Model.xlsx, Forecast and DCF tabs (base case)"),
        ("FactFinancials", "Adj. EBITDA, net debt / adj. EBITDA, DSO, backlog", "; ".join(f"{y}: {s}" for y, s in REPORTED_SRC.items())),
        ("FactSegment", "Net revenue by reportable segment", f"2023: {SRC_MDA23}; 2024-2025: {SRC_FY25}"),
        ("Valuation, Sensitivity, Comps, EVBridge", "DCF, WACC and peer multiples", "Stantec_DCF_Model.xlsx, DCF / WACC / Comps tabs"),
    ], columns=["Table", "Fields", "Source"])
    xl = ROOT / "powerbi" / "Stantec_Dashboard_Data.xlsx"
    with pd.ExcelWriter(xl, engine="openpyxl") as xw:
        for name, df in [("FactFinancials", fact), ("FactSegment", seg), ("DimYear", years), ("Valuation", valuation),
                         ("Sensitivity", sens), ("Comps", comps), ("EVBridge", bridge), ("Sources", sources)]:
            df.to_excel(xw, sheet_name=name, index=False)
    print(f"Wrote {xl.name}: {len(fact)} years, {len(seg)} segment rows, {len(sens)} sensitivity cells")


if __name__ == "__main__":
    main()
